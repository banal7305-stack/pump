"""F6 이상진단 보고서 (.md / .docx / 우선순위표 .xlsx) 와 종합 의견(템플릿 또는 Claude API)."""
from __future__ import annotations

import json
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd

from .loader import DataContext, ROOT
from .models import LABELS, UNITS, DiagnosisReport

log = logging.getLogger(__name__)
Block = tuple  # ("h1"|"h2"|"p"|"list"|"olist"|"table", payload)


# ---------------- 종합 의견 ----------------
def _fmt_days(d: Optional[float]) -> str:
    return "이미 초과" if d == 0 else f"약 {d:.0f}일 뒤"


def facts(rep: DiagnosisReport) -> dict:
    """LLM에 넘기는 계산 결과 요약 (LLM은 이 값만 근거로 쓴다)."""
    det, pr = rep.detection, rep.priority
    return {
        "펌프": f"{rep.pump.pump_id} {rep.pump.name}",
        "측정시각": f"{rep.measured_at:%Y-%m-%d %H:%M}",
        "종합상태": det.status,
        "항목판정": [{"항목": g.label, "값": g.value, "단위": UNITS[g.variable],
                   "정격대비%": None if g.ratio is None else round(g.ratio * 100),
                   "등급": g.grade_name} for g in det.grades if g.value is not None],
        "추세": {LABELS[v]: round(s, 3) for v, s in det.trend.slopes.items()} if det.trend.available else det.trend.note,
        "경고도달예상일": {LABELS[v]: (None if e is None else round(e)) for v, e in det.trend.eta_days.items()},
        "현장징후": det.field_signs,
        "원인후보": [{"순위": i + 1, "원인": c.mode, "신뢰도%": round(c.confidence), "신뢰도": c.confidence_label,
                   "근거": c.evidence[:4]} for i, c in enumerate(rep.candidates)],
        "추가확인": [c.item for c in rep.checks[:4]],
        "우선순위": {"점수": pr.score, "등급": pr.level, "기한": pr.deadline, "작업순서": pr.work_steps},
        "오경보메모": det.false_alarm_note,
    }


def template_opinion(rep: DiagnosisReport) -> str:
    det, pr = rep.detection, rep.priority
    if det.stopped:
        return "펌프가 정지 상태(유량·전류 거의 0)로 입력되어 운전 이상 진단을 하지 않았습니다. 운전 중 값을 입력하세요."
    bad = [g for g in det.grades if g.grade >= 1]
    parts = [f"{rep.pump.pump_id}의 종합 상태는 '{det.status}'입니다."]
    if bad:
        parts.append(", ".join(f"{g.label} {g.value:g} {UNITS[g.variable]}({g.grade_name})" for g in bad) + "이(가) 기준을 벗어났습니다.")
    if det.status == "정상":
        parts.append("평소 운전 범위 안에 있어 즉시 조치할 사항은 없습니다. 정기 점검 주기를 유지하세요.")
        return " ".join(parts)
    if rep.candidates:
        c = rep.candidates[0]
        parts.append(f"가장 가능성이 높은 원인은 {c.mode}(신뢰도 {c.confidence_label}, {c.confidence:.0f}%)입니다.")
    etas = [(LABELS[v], e) for v, e in det.trend.eta_days.items() if e is not None and e <= 30]
    if etas:
        parts.append("현재 추세면 " + ", ".join(f"{n} 경고 기준에 {_fmt_days(e)}" for n, e in etas) + " 도달합니다.")
    parts.append(f"우선순위는 '{pr.level}'({pr.deadline})입니다.")
    if rep.checks:
        parts.append(f"먼저 '{rep.checks[0].item}'부터 확인하세요.")
    if det.false_alarm_note:
        parts.append(det.false_alarm_note)
    return " ".join(parts)


# ---------------- LLM 응답 검증 ----------------
# 식별자·날짜는 숫자로 보지 않고 토큰 그대로 비교한다 (P-105의 105, 2026-09-30의 30 등)
_TOKEN_RE = re.compile(r"[A-Za-z]+-[A-Za-z]*\d+|\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2})?|\d{2}-\d{2} \d{2}:\d{2}|\d{1,2}:\d{2}")
_NUM_RE = re.compile(r"(?<![\w.])([-+]?)(\d+(?:\.\d+)?)")
_IDENT_KEYS = {"펌프", "측정시각"}


def _split(text: str) -> tuple[set[str], list[tuple[float, int, bool]]]:
    """텍스트 → (식별자·날짜 토큰, [(값, 소수 자릿수, 명시적 음수 여부)])."""
    tokens = set(_TOKEN_RE.findall(text))
    rest = _TOKEN_RE.sub(" ", text)
    nums = []
    for sign, n in _NUM_RE.findall(rest):
        decimals = len(n.split(".")[1]) if "." in n else 0
        value = float(n)
        nums.append((-value if sign == "-" else value, decimals, sign == "-"))
    return tokens, nums


def _allowed(obj, top_n: int = 3) -> tuple[set[str], list[float]]:
    """계산 결과(facts)에서 쓸 수 있는 토큰과 숫자. 순위 같은 작은 정수(1~top_n)는 항상 허용."""
    tokens: set[str] = set()
    values: list[float] = [float(i) for i in range(1, top_n + 1)]

    def walk(o, key=None):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, k)
        elif isinstance(o, (list, tuple)):
            for v in o:
                walk(v, key)
        elif isinstance(o, bool) or o is None:
            return
        elif isinstance(o, (int, float)):
            values.append(float(o))
        elif isinstance(o, str):
            t, nums = _split(o)
            tokens.update(t)
            if key not in _IDENT_KEYS:
                values.extend(v for v, _, _ in nums)
    walk(obj)
    return tokens, values


def unsupported_numbers(text: str, data: dict) -> list[str]:
    """LLM 문장에서 계산 결과로 뒷받침되지 않는 숫자·식별자 목록 (반올림·부호 고려)."""
    ok_tokens, ok_values = _allowed(data)
    tokens, nums = _split(text)
    bad = sorted(tokens - ok_tokens)
    for value, decimals, explicit_neg in nums:
        def matches(a: float) -> bool:
            target = a if explicit_neg else abs(a)
            return abs(round(target, decimals) - (value if explicit_neg else abs(value))) < 1e-9
        if not any(matches(a) for a in ok_values):
            bad.append(f"{value:g}")
    return bad


def semantic_problems(text: str, data: dict) -> list[str]:
    """상태·1순위 원인·우선순위 등급이 계산 결과와 다르게 쓰였는지 확인."""
    problems = []
    status, level = data["종합상태"], data["우선순위"]["등급"]
    for said in re.findall(r"상태[는가은이]?\s*['\"‘“]?(정상|이상징후|이상)", text):
        if said != status:
            problems.append(f"상태 '{said}'(계산 '{status}')")
    for said in re.findall(r"우선순위[는가은이]?\s*['\"‘“]?(즉시|긴급|계획|관찰)", text):
        if said != level:
            problems.append(f"우선순위 '{said}'(계산 '{level}')")
    if status not in text:
        problems.append("종합 상태 누락")
    if level not in text:
        problems.append("우선순위 등급 누락")
    if data["원인후보"]:
        top = data["원인후보"][0]["원인"]
        core = re.split(r"[(/]", top)[0].strip()
        if core not in text:
            problems.append(f"1순위 원인 '{top}' 누락")
    return problems


SYSTEM_PROMPT = (
    "당신은 제철소 펌프 설비 보전 엔지니어를 돕는 진단 보고서 작성자입니다. "
    "사용자가 주는 JSON은 규칙 엔진이 계산한 진단 결과입니다. 이 결과만 근거로 보고서의 '종합 의견'을 한국어 한 문단(4~6문장)으로 쓰세요. "
    "JSON에 없는 숫자·원인·부품을 만들지 말고, 상태·원인 순위·우선순위 등급을 바꾸지 마세요. "
    "현장 담당자가 바로 행동할 수 있게 무엇이 이상한지, 가장 유력한 원인과 근거, 먼저 확인할 일, 조치 기한 순서로 쓰세요. "
    "제목·목록·마크다운 없이 문장만 쓰세요."
)


def llm_opinion(ctx: DataContext, rep: DiagnosisReport) -> Optional[str]:
    """Claude API로 종합 의견 생성. 인증 정보가 없거나 실패·검증 실패 시 None."""
    lcfg = ctx.cfg["llm"]
    if not lcfg.get("enabled") or not (os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")):
        return None
    try:
        import anthropic
    except ImportError:
        return None
    data = facts(rep)
    client = anthropic.Anthropic(timeout=float(lcfg["timeout_sec"]), max_retries=2)
    try:
        resp = client.beta.messages.create(
            model=lcfg["model"],
            max_tokens=int(lcfg["max_tokens"]),
            output_config={"effort": lcfg["effort"]},
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": json.dumps(data, ensure_ascii=False, default=str)}],
        )
    except anthropic.AuthenticationError:
        log.warning("Claude API 인증 실패: 템플릿 의견을 사용합니다.")
        return None
    except anthropic.RateLimitError:
        log.warning("Claude API 요청 한도 초과: 템플릿 의견을 사용합니다.")
        return None
    except anthropic.APIStatusError as e:
        log.warning("Claude API 오류(%s): 템플릿 의견을 사용합니다.", e.status_code)
        return None
    except anthropic.APIConnectionError:
        log.warning("Claude API 연결 실패: 템플릿 의견을 사용합니다.")
        return None
    if resp.stop_reason != "end_turn":     # 거절(refusal)·잘림(max_tokens) 등은 쓰지 않는다
        log.warning("Claude 응답 종료 사유 %s: 템플릿 의견을 사용합니다.", resp.stop_reason)
        return None
    text = " ".join(b.text for b in resp.content if b.type == "text").strip()
    if not text:
        return None
    unknown = unsupported_numbers(text, data)
    problems = semantic_problems(text, data)
    if unknown or problems:
        log.warning("LLM 의견 검증 실패(숫자 %s, 내용 %s): 템플릿 의견을 사용합니다.", unknown, problems)
        return None
    return text


def write_opinion(ctx: DataContext, rep: DiagnosisReport, use_llm: bool = True) -> None:
    text = llm_opinion(ctx, rep) if use_llm and not rep.detection.stopped else None
    rep.opinion, rep.opinion_source = (text, "llm") if text else (template_opinion(rep), "template")


# ---------------- 보고서 블록 ----------------
def _num(v, nd=2) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    return f"{v:,.{nd}f}".rstrip("0").rstrip(".") if isinstance(v, float) else f"{v:,}"


def blocks(ctx: DataContext, rep: DiagnosisReport) -> list[Block]:
    p, det, pr = rep.pump, rep.detection, rep.priority
    b: list[Block] = [("h1", f"펌프 이상진단 보고서 — {p.pump_id} {p.name}"),
                      ("p", f"측정 시각 {rep.measured_at:%Y-%m-%d %H:%M} · 보고서 생성 {datetime.now():%Y-%m-%d %H:%M} · 데이터 {ctx.source.name}")]
    b += [("h2", "1. 설비 정보"), ("table", (["항목", "내용"], [
        ["펌프ID / 설비명", f"{p.pump_id} / {p.name}"], ["위치 / 형식", f"{p.location} / {p.pump_type}"],
        ["모터", f"{p.kw:g} kW"],
        ["정격", f"유량 {p.rated['flow']:g} m³/h, 토출압력 {p.rated['discharge']:g} bar, 흡입압력 {p.rated['suction']:g} bar, 전류 {p.rated['current']:g} A"],
        ["설치일", f"{p.installed:%Y-%m-%d}" if p.installed is not None else "-"], ["설비중요도", f"{p.importance:g}"]]))]

    rows = []
    for g in det.grades:
        rows.append([g.label, "미입력" if g.value is None else f"{g.value:g} {UNITS[g.variable]}",
                     "-" if g.ratio is None else f"{g.ratio * 100:.0f}%",
                     "-" if g.baseline is None else f"{g.baseline:.3g}",
                     "-" if g.z is None else f"{g.z:+.1f}", g.grade_name])
    b += [("h2", "2. 입력값과 판정"), ("table", (["항목", "값", "정격 대비", "평소 기준선", "z-score", "등급"], rows))]

    b += [("h2", "3. 종합 상태"),
          ("p", f"종합 상태: **{det.status}** (최대 등급 {['정상', '관찰', '주의', '경고'][max(0, det.max_grade)]}, 이상지수 {det.anomaly_index:.1f})"),
          ("p", f"종합 의견{' (Claude 작성)' if rep.opinion_source == 'llm' else ''}: {rep.opinion}")]
    if det.false_alarm_note:
        b.append(("p", f"주의: {det.false_alarm_note}"))

    b.append(("h2", "4. 추세·알람 요약"))
    if det.trend.available:
        tl = [f"최근 {ctx.cfg['trend']['window_days']}일 운전 {det.trend.points}건 기준, 진행 속도 점수 {det.trend.speed_score:.0f}"]
        for v in ("vibration", "bearing_temp", "current", "flow", "discharge", "suction"):
            if v in det.trend.slopes:
                s = det.trend.slopes[v]
                unit = UNITS[v] if v in ("vibration", "bearing_temp") else "%p(정격 대비)"
                val = s if v in ("vibration", "bearing_temp") else s * 100
                tl.append(f"{LABELS[v]}: 일 {val:+.2f} {unit}")
        for v, e in det.trend.eta_days.items():
            if e is not None:
                tl.append(f"{LABELS[v]} 경고 기준 도달: {_fmt_days(e)}")
        b.append(("list", tl))
    else:
        b.append(("p", det.trend.note))
    if det.alarms:
        b.append(("table", (["알람코드", "알람명", f"최근 {ctx.cfg['alarms']['context_days']}일 건수", "연속 발생(일)", "최근 값", "추세", "오경보 가능"],
                            [[a.code, a.name, a.count, a.consecutive_days, ", ".join(f"{v:g}" for v in a.values[-3:]) or "-",
                              "상승" if a.rising else "-", a.suspected_false or "-"] for a in det.alarms])))
    else:
        b.append(("p", f"최근 {ctx.cfg['alarms']['context_days']}일 알람 없음"))
    if det.field_signs:
        b.append(("p", "현장 징후: " + ", ".join(f"{s} ({det.sign_sources[s]})" for s in det.field_signs)))

    b.append(("h2", "5. 원인 후보"))
    if rep.candidates:
        b.append(("table", (["순위", "원인", "신뢰도", "근거 수"],
                            [[i + 1, c.mode, f"{c.confidence_label} ({c.confidence:.0f}%)", len(c.evidence)] for i, c in enumerate(rep.candidates)])))
        for i, c in enumerate(rep.candidates):
            items = [f"근거: {e}" for e in c.evidence] + [f"불일치: {m}" for m in c.mismatches] + [f"참고: {n}" for n in c.notes]
            b += [("p", f"**{i + 1}순위 {c.mode}**"), ("list", items or ["근거 없음"])]
    else:
        b.append(("p", "이상이 없어 원인 후보를 제시하지 않습니다." if det.status == "정상" else "원인 후보를 찾지 못했습니다."))

    b += [("h2", "6. 추가 확인 항목"),
          ("olist", [f"{c.item} — {c.reason}" for c in rep.checks] or ["추가 확인 항목 없음"])]

    comp = pr.components
    b += [("h2", "7. 점검·교체 우선순위"),
          ("p", f"우선순위: **{pr.level}** ({pr.deadline}) · 점수 {pr.score:.1f}"),
          ("table", (["구성", "점수(0~100)", "가중치"],
                     [[n, f"{comp[k]:.0f}", f"{ctx.cfg['priority']['weights'][k]:.2f}"]
                      for k, n in (("severity", "심각도"), ("speed", "진행 속도"), ("importance", "설비 중요도"), ("confidence", "원인 신뢰도"))]))]
    if pr.work_steps:
        b.append(("olist", pr.work_steps))
    if pr.case_count:
        b.append(("p", f"교체·수리 대상(과거 사례): {pr.parts} · 평균 작업시간 {pr.avg_hours:.1f}시간 · 평균 비용 {pr.avg_cost:,.0f}원 (사례 {pr.case_count}건)"))

    b.append(("h2", "8. 유사 정비 사례"))
    seen, rows = set(), []
    for c in rep.candidates:
        for s in c.similar_cases:
            if s.wo_id in seen:
                continue
            seen.add(s.wo_id)
            rows.append([s.wo_id, f"{s.date:%Y-%m-%d}", s.pump_id, c.mode, s.cause, s.symptom, s.action, f"{s.hours:g}h / {s.cost:,}원"])
    b.append(("table", (["작업ID", "일자", "펌프", "관련 원인", "확인된 고장원인", "작업 전 증상", "조치", "시간/비용"], rows))
             if rows else ("p", "유사 사례 없음"))
    b.append(("p", f"※ {ctx.cfg['report']['disclaimer']}"))
    return b


def to_markdown(bl: list[Block]) -> str:
    out = []
    for kind, payload in bl:
        if kind == "h1":
            out.append(f"# {payload}\n")
        elif kind == "h2":
            out.append(f"\n## {payload}\n")
        elif kind == "p":
            out.append(f"{payload}\n")
        elif kind in ("list", "olist"):
            out += [f"{i + 1}. {x}" if kind == "olist" else f"- {x}" for i, x in enumerate(payload)]
            out.append("")
        elif kind == "table":
            head, rows = payload
            cell = lambda x: str(x).replace("|", "/").replace("\n", " ")
            out.append("| " + " | ".join(head) + " |")
            out.append("|" + "---|" * len(head))
            out += ["| " + " | ".join(cell(x) for x in r) + " |" for r in rows]
            out.append("")
    return "\n".join(out)


def to_docx(bl: list[Block], path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    def para(text: str, style_name: Optional[str] = None):
        p = doc.add_paragraph(style=style_name)
        for i, chunk in enumerate(text.split("**")):
            if chunk:
                p.add_run(chunk).bold = i % 2 == 1
    for kind, payload in bl:
        if kind == "h1":
            doc.add_heading(payload, level=0)
        elif kind == "h2":
            doc.add_heading(payload, level=1)
        elif kind == "p":
            para(payload)
        elif kind in ("list", "olist"):
            for x in payload:
                para(x, "List Number" if kind == "olist" else "List Bullet")
        elif kind == "table":
            head, rows = payload
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Light Grid Accent 1"
            for c, h in zip(t.rows[0].cells, head):
                c.text = str(h)
            for r in rows:
                for c, x in zip(t.add_row().cells, r):
                    c.text = str(x)
    doc.save(path)


def priority_rows(reports: list[DiagnosisReport]) -> pd.DataFrame:
    rows = []
    for r in reports:
        top = r.candidates[0] if r.candidates else None
        rows.append({"펌프ID": r.pump.pump_id, "설비명": r.pump.name, "측정시각": r.measured_at,
                     "종합상태": r.detection.status, "1순위 원인": top.mode if top else "-",
                     "신뢰도(%)": round(top.confidence) if top else None,
                     "우선순위 점수": r.priority.score, "등급": r.priority.level, "기한": r.priority.deadline,
                     "심각도": round(r.priority.components["severity"]), "진행속도": round(r.priority.components["speed"]),
                     "설비중요도": round(r.priority.components["importance"]),
                     "첫 확인 항목": r.checks[0].item if r.checks else "-"})
    return pd.DataFrame(rows).sort_values("우선순위 점수", ascending=False).reset_index(drop=True)


def to_xlsx(reports: list[DiagnosisReport], path: Path) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment

    df = priority_rows(reports)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="우선순위", index=False)
        ws = xw.sheets["우선순위"]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, w in zip("ABCDEFGHIJKLM", [9, 20, 17, 10, 22, 10, 12, 8, 14, 8, 9, 10, 40]):
            ws.column_dimensions[col].width = w
        for (c,) in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            c.number_format = "yyyy-mm-dd hh:mm"
        ws.freeze_panes = "A2"


def output_dir(ctx: DataContext) -> Path:
    d = Path(ctx.cfg["data"]["output_dir"])
    d = d if d.is_absolute() else ROOT / d
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_report(ctx: DataContext, rep: DiagnosisReport) -> dict[str, Path]:
    bl = blocks(ctx, rep)
    stem = f"report_{rep.pump.pump_id}_{rep.measured_at:%Y%m%d_%H%M}"
    d = output_dir(ctx)
    paths = {"md": d / f"{stem}.md", "docx": d / f"{stem}.docx", "xlsx": d / f"{stem}_priority.xlsx"}
    paths["md"].write_text(to_markdown(bl), encoding="utf-8")
    to_docx(bl, paths["docx"])
    to_xlsx([rep], paths["xlsx"])
    return paths
